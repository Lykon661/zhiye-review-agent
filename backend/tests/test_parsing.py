from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from app.parsing import parse_document


def make_image(path: Path) -> Path:
    Image.new("RGB", (80, 50), "white").save(path)
    return path


def test_pdf_page_ordinals_include_blank_pages(tmp_path: Path) -> None:
    path = tmp_path / "lecture.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(30, 760, "Chapter one: Cache locality")
    pdf.showPage()
    pdf.showPage()
    pdf.drawString(30, 760, "Chapter two: Page replacement")
    pdf.drawString(30, 735, "Algorithm      Complexity")
    pdf.drawString(30, 710, "FIFO           O(1)")
    pdf.save()

    result = parse_document(path, path.name)

    assert (result.units_total, result.units_with_text) == (3, 2)
    assert [block.locator["page"] for block in result.blocks] == [1, 3]
    assert [block.order for block in result.blocks] == [0, 1]
    assert "Page replacement" in result.blocks[-1].text
    assert "FIFO" in result.blocks[-1].text
    assert result.warnings == [], "Partial recognition diagnostics stay out of the public parse result."


def test_pdf_skips_unreadable_images_without_public_warnings(tmp_path: Path) -> None:
    image = make_image(tmp_path / "scan.png")
    mixed = tmp_path / "mixed.pdf"
    pdf = canvas.Canvas(str(mixed))
    pdf.drawString(30, 760, "Visible digital text")
    pdf.drawImage(str(image), 30, 600, width=80, height=50)
    pdf.save()
    result = parse_document(mixed, mixed.name)
    assert result.warnings == []
    assert "Visible digital text" in result.blocks[0].text

    scan = tmp_path / "scan.pdf"
    pdf = canvas.Canvas(str(scan))
    pdf.drawImage(str(image), 30, 600, width=80, height=50)
    pdf.save()
    with pytest.raises(ValueError, match="未提取到可用内容"):
        parse_document(scan, scan.name)


def test_password_protected_pdf_is_rejected(tmp_path: Path) -> None:
    original = tmp_path / "original.pdf"
    pdf = canvas.Canvas(str(original))
    pdf.drawString(30, 760, "Confidential course")
    pdf.save()
    encrypted = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.append(PdfReader(str(original)))
    writer.encrypt("secret")
    writer.write(str(encrypted))

    with pytest.raises(ValueError, match="PDF 需要密码|PDF 已加密"):
        parse_document(encrypted, encrypted.name)


def test_pptx_group_text_table_and_slide_positions(tmp_path: Path) -> None:
    path = tmp_path / "slides.pptx"
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[5])
    first.shapes.title.text = "第一章 内存管理"
    group = first.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "组内文本：虚拟内存"
    nested = group.shapes.add_group_shape()
    nested.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1)).text = "深层组合：分页"
    table = first.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "算法"
    table.cell(0, 1).text = "条件"
    table.cell(1, 0).text = "LRU"
    table.cell(1, 1).text = "局部性"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "第二张正文"
    second.shapes.add_picture(str(make_image(tmp_path / "picture.png")), Inches(1), Inches(3))
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(str(path))

    result = parse_document(path, path.name)

    assert (result.units_total, result.units_with_text) == (3, 2)
    assert [block.order for block in result.blocks] == list(range(len(result.blocks)))
    assert [block.text for block in result.blocks[:3]] == ["第一章 内存管理", "组内文本：虚拟内存", "深层组合：分页"]
    assert result.blocks[2].locator["shape_path"].count(".") == 2
    table_block = next(block for block in result.blocks if block.block_type == "table")
    assert table_block.text == "算法\t条件\nLRU\t局部性"
    assert table_block.heading == "第一章 内存管理"
    assert result.blocks[-1].locator["slide"] == 2
    assert result.warnings == []


def test_docx_preserves_body_order_heading_and_table_position(tmp_path: Path) -> None:
    path = tmp_path / "notes.docx"
    document = Document()
    document.add_heading("第一章 内存管理", level=1)
    intro = document.add_paragraph("正文开始：地址转换。")
    intro._p.append(OxmlElement("m:oMath"))
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "术语"
    table.cell(0, 1).text = "定义"
    table.cell(1, 0).text = "页"
    table.cell(1, 1).text = "固定长度块"
    document.add_paragraph("表格之后的正文。")
    document.add_heading("第二章 调度", level=1)
    document.add_paragraph("正文结束：时间片。")
    document.add_picture(str(make_image(tmp_path / "inline.png")))
    document.save(str(path))

    result = parse_document(path, path.name)

    assert (result.units_total, result.units_with_text) == (8, 6)
    assert result.blocks[0].block_type == "heading"
    assert result.blocks[1].locator["paragraph"] == 2
    assert result.blocks[2].block_type == "table"
    assert result.blocks[2].locator["table"] == 1
    assert "页\t固定长度块" in result.blocks[2].text
    assert result.blocks[3].text == "表格之后的正文。"
    assert result.blocks[3].locator["paragraph"] == 4
    assert result.blocks[-1].heading == "第二章 调度"
    assert all("page" not in block.locator for block in result.blocks)
    assert result.warnings == []


def test_docx_nested_table_text_is_not_silently_lost(tmp_path: Path) -> None:
    path = tmp_path / "nested.docx"
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "嵌套前"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套表格内容"
    cell.add_paragraph("嵌套后")
    document.save(str(path))

    result = parse_document(path, path.name)
    assert result.blocks[0].text == "嵌套前\n嵌套表格内容\n嵌套后"
    assert result.blocks[0].locator["table"] == 1


def test_pptx_soft_line_break_stays_a_word_boundary(tmp_path: Path) -> None:
    path = tmp_path / "softbreak.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    textbox.text = "第一行\v第二行"
    deck.save(str(path))
    result = parse_document(path, path.name)
    assert result.blocks[0].text == "第一行\n第二行"


@pytest.mark.parametrize("suffix", [".doc", ".ppt", ".rtf"])
def test_legacy_formats_report_missing_tika(tmp_path: Path, suffix: str, monkeypatch) -> None:
    path = tmp_path / f"old{suffix}"
    path.write_bytes(b"old")
    monkeypatch.setenv("TIKA_APP_JAR", str(tmp_path / "missing-tika.jar"))
    with pytest.raises(ValueError, match="Tika 解析组件尚未安装"):
        parse_document(path, path.name)


@pytest.mark.parametrize("suffix", [".pdf", ".pptx", ".docx"])
def test_damaged_documents_have_readable_errors(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / f"damaged{suffix}"
    path.write_bytes(b"not a valid document")
    with pytest.raises(ValueError, match="文件解析失败"):
        parse_document(path, path.name)


def test_empty_and_unknown_files_are_rejected(tmp_path: Path) -> None:
    path = tmp_path / "empty.pdf"
    path.touch()
    with pytest.raises(ValueError, match="资料文件为空"):
        parse_document(path, path.name)
    with pytest.raises(ValueError, match="不支持该文件格式"):
        parse_document(path, "data.xlsx")


@pytest.mark.parametrize("suffix", [".pptx", ".docx"])
def test_textless_office_documents_are_rejected(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / f"empty{suffix}"
    if suffix == ".pptx":
        document = Presentation()
        document.slides.add_slide(document.slide_layouts[6])
    else:
        document = Document()
    document.save(str(path))
    with pytest.raises(ValueError, match="未提取到可用内容"):
        parse_document(path, path.name)
