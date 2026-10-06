"""Controlled model doubles; parsing, SQLite, Chroma and HTTP remain real."""
from __future__ import annotations

from collections import deque
from hashlib import sha256
from io import BytesIO
import json
import math
import re

from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
import pytest
from reportlab.pdfgen.canvas import Canvas

from app.config import Settings


class DeterministicEmbedder:
    """Stable hashed words/ngrams for offline integration tests, not retrieval quality."""
    tokenizer = None

    def __init__(self):
        self.query_calls: list[list[str]] = []

    def embed(self, texts: list[str], query: bool = False) -> list[list[float]]:
        if query:
            self.query_calls.append(list(texts))
        vectors = []
        for text in texts:
            tokens = re.findall(r"[a-z0-9]+|[\u4e00-\u9fff]{1,}", text.casefold())
            words = []
            for token in tokens:
                if re.search(r"[\u4e00-\u9fff]", token):
                    words.extend(token[i:i + 2] for i in range(max(1, len(token) - 1)))
                else:
                    words.append(token)
            vector = [0.0] * 128
            for word in words:
                vector[int.from_bytes(sha256(word.encode()).digest()[:4], "big") % len(vector)] += 1
            norm = math.sqrt(sum(value * value for value in vector)) or 1
            vectors.append([value / norm for value in vector])
        return vectors


class ControlledLLM:
    configured = True

    def __init__(self):
        self.complete_calls: list[list[dict]] = []
        self.stream_calls: list[list[dict]] = []
        self.complete_responses = deque()
        self.stream_responses = deque()
        self.fail_complete = False
        self.fail_stream = False

    def complete(self, messages: list[dict], json_mode: bool = False) -> str:
        self.complete_calls.append(messages)
        if self.fail_complete:
            raise RuntimeError("Controlled LLM timeout")
        if self.complete_responses:
            value = self.complete_responses.popleft()
            if isinstance(value, Exception):
                raise value
            return value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
        prompt = "\n".join(str(message.get("content", "")) for message in messages)
        payload = json.loads(messages[-1]["content"])
        if "[TASK:PLAN]" in prompt:
            question = payload["question"]
            tool = "find_past_questions" if ("真题" in question or "past question" in question.lower()) else "search_materials"
            if payload.get("observations"):
                return json.dumps({"query": question, "tools": [], "done": True}, ensure_ascii=False)
            # Resolve a simple pronoun using history so the application can be checked.
            query = "RAG" if "它" in question else question
            return json.dumps({"query": query, "tools": [{"name": tool, "query": query}]}, ensure_ascii=False)
        if "[TASK:OUTLINE_MAP]" in prompt:
            points = [{"text": block["text"], "source_ids": [block["source_id"]]} for block in payload["blocks"] if block["text"].strip()]
            return json.dumps({"core_points": points, "definitions": [], "confusions": [], "past_questions": []}, ensure_ascii=False)
        if "[TASK:OUTLINE_MERGE]" in prompt:
            result = {key: [] for key in ("core_points", "definitions", "confusions", "past_questions")}
            for partial in payload["partials"]:
                for key in result:
                    result[key].extend(partial.get(key, []))
            return json.dumps(result, ensure_ascii=False)
        raise AssertionError(f"Unknown LLM task: {prompt[:200]}")

    def stream(self, messages: list[dict]):
        self.stream_calls.append(messages)
        if self.fail_stream:
            raise RuntimeError("Controlled LLM stream timeout")
        if self.stream_responses:
            value = self.stream_responses.popleft()
        else:
            prompt = "\n".join(str(message.get("content", "")) for message in messages)
            ids = re.findall(r"S[a-f0-9]{32}", prompt)
            value = f"RAG 根据检索证据生成回答。[{ids[0]}]" if ids else "当前资料不足，无法根据复习资料回答这个问题。"
        if isinstance(value, Exception):
            raise value
        for part in (value[:len(value) // 2], value[len(value) // 2:]):
            if part:
                yield part


def make_docx(paragraphs: list[str] | None = None, include_table: bool = True) -> bytes:
    document = Document()
    document.add_heading("第一章 RAG", level=1)
    for text in paragraphs or ["RAG combines retrieval with generation.", "A knowledge base preserves source evidence."]:
        document.add_paragraph(text)
    if include_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Concept"
        table.cell(0, 1).text = "Definition"
        table.cell(1, 0).text = "Retrieval"
        table.cell(1, 1).text = "Find relevant evidence."
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_pdf(pages: list[str] | None = None) -> bytes:
    output = BytesIO()
    canvas = Canvas(output)
    for text in pages or ["RAG retrieval uses source evidence.", "Generation follows the retrieved facts."]:
        canvas.drawString(40, 790, text)
        canvas.showPage()
    canvas.save()
    return output.getvalue()


def make_pptx() -> bytes:
    presentation = Presentation()
    for number, text in enumerate(["RAG retrieval source evidence", "RAG generation grounding"], 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
        textbox.text_frame.text = f"第{number}章 {text}"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(7), Inches(1)).table
        table.cell(0, 0).text = "Concept"
        table.cell(1, 0).text = "Evidence"
        table.cell(1, 1).text = "Original source"
        group = slide.shapes.add_group_shape()
        child = group.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
        child.text_frame.text = "Grouped text is also evidence."
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


@pytest.fixture
def llm():
    return ControlledLLM()


@pytest.fixture
def embedder():
    return DeterministicEmbedder()


@pytest.fixture
def settings(tmp_path):
    return Settings(data_dir=tmp_path / "data", retrieval_min_score=0.05)


@pytest.fixture
def app(settings, embedder, llm):
    from app.main import create_app
    return create_app(settings=settings, embedder=embedder, llm=llm)


@pytest.fixture
def client(app):
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture
def course(client):
    response = client.post("/courses", json={"name": "RAG 期末复习"})
    assert response.status_code in (200, 201), response.text
    return response.json()
