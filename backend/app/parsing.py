"""Document text, local image OCR and Office formulas with source positions."""

from __future__ import annotations

from dataclasses import dataclass, field
from html import escape
import json
import logging
from pathlib import Path
import re
import unicodedata
from typing import Any, Iterator
from urllib.parse import unquote
from xml.etree import ElementTree

from .office_math import office_formulas
from .legacy_wmf_math import wmf_formula
from .question_images import save_question_image
from .tika import TikaError, extracted_document
from .upload_formats import IMAGE_SUFFIXES, SUPPORTED_FORMATS_LABEL, SUPPORTED_SUFFIXES, TEXT_SUFFIXES, TIKA_SUFFIXES

logger = logging.getLogger(__name__)


@dataclass
class ParsedBlock:
    text: str
    locator: dict[str, Any]
    heading: str | None = None
    block_type: str = "text"
    order: int = 0


@dataclass
class ParseResult:
    blocks: list[ParsedBlock] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    units_total: int = 0
    units_with_text: int = 0
    images_processed: int = 0
    ocr_blocks: int = 0
    formula_blocks: int = 0
    image_dir: Path | None = field(default=None, repr=False)


def _clean_text(text: str | None) -> str:
    # Preserve line breaks and table tabs; drop XML/PDF control characters only.
    normalized = (text or "").replace("\r\n", "\n").replace("\r", "\n").replace("\v", "\n")
    return re.sub(r"[\x00-\x08\x0c\x0e-\x1f]", "", normalized).strip()


def _script_text(text: str, kind: str | None) -> str:
    # Retain source formatting without guessing whether a trailing digit is a power.
    return f"<{kind}>{escape(text, quote=False)}</{kind}>" if text and kind else text


def _pptx_frame_text(frame: Any) -> str:
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    paragraphs = []
    for paragraph in frame.paragraphs:
        default = paragraph._p.find(f"{ns}pPr/{ns}defRPr")
        parts = []
        for element in paragraph._p:
            tag = element.tag.rsplit("}", 1)[-1]
            if tag == "br":
                parts.append("\n")
            elif tag in {"r", "fld"}:
                value = element.find(f"{ns}t")
                if value is None:
                    continue
                properties = element.find(f"{ns}rPr")
                baseline = properties.get("baseline") if properties is not None else None
                if baseline is None and default is not None:
                    baseline = default.get("baseline")
                offset = int(baseline or "0")
                parts.append(_script_text(value.text or "", "sup" if offset > 0 else "sub" if offset < 0 else None))
        paragraphs.append("".join(parts))
    return "\n".join(paragraphs)


def _docx_paragraph_text(paragraph: Any) -> str:
    def script_kind(run):
        # A Word run can inherit vertical alignment from its character/paragraph style.
        fonts = [run.font]
        for style in (run.style, paragraph.style):
            seen = set()
            while style is not None and style.style_id not in seen:
                seen.add(style.style_id)
                fonts.append(style.font)
                style = style.base_style
        for font in fonts:
            if font.superscript is not None or font.subscript is not None:
                return "sup" if font.superscript else "sub" if font.subscript else None
        return None

    parts = []
    # iter_inner_content keeps hyperlinks at their original position as well.
    for item in paragraph.iter_inner_content():
        for run in getattr(item, "runs", (item,)):
            parts.append(_script_text(run.text, script_kind(run)))
    return "".join(parts)


def _add_office_math(result, element, locator, heading):
    added = False
    for index, formula in enumerate(office_formulas(element), 1):
        if formula["unsupported"]:
            logger.debug("Skipping unsupported Office formula at %s", locator["label"])
            continue
        added |= _add_block(result, f"[Office 原生公式]\n$$\n{formula['latex']}\n$$", {
            **locator, "label": locator["label"] + f" · 公式 {index}",
            "recognition": "office_math", "formula": index,
        }, heading, "formula")
        result.formula_blocks += 1
    return added


def _add_visual(result, visual, image, locator, heading=None, context="", existing_text=""):
    added = False
    if result.image_dir is not None:
        try:
            url = save_question_image(image, result.image_dir)
            if url:
                added = _add_block(result, f"![{locator['label']} · 原图]({url})",
                                   {**locator, "image_url": url}, heading, "image")
        except Exception:
            logger.debug("Unable to preserve question figure at %s", locator["label"], exc_info=True)
    if visual is None or not visual.settings.ocr_enabled:
        result.warnings.append(f"{locator['label']}含图片，OCR 未启用。")
        return added
    try:
        recognized = visual.recognize(image, context=context or heading or "", existing_text=existing_text)
    except Exception:
        if not added:
            raise
        logger.debug("OCR failed; original question figure retained", exc_info=True)
        return added
    if getattr(recognized, "skipped", False):
        return added
    result.images_processed += 1
    if recognized.lines:
        text = "[图片文字 OCR 自动识别]\n" + "\n".join(line["text"] for line in recognized.lines)
        added |= _add_block(result, text, {**locator, "recognition": "ocr", "ocr_lines": recognized.lines,
                                           "label": locator["label"] + " · 图片文字 OCR"}, heading, "ocr")
        result.ocr_blocks += 1
    for index, formula in enumerate(recognized.formulas, 1):
        text = "[图片公式自动识别]\n$$\n" + formula["latex"] + "\n$$"
        if formula.get("ocr_text"):
            text += "\n对应图片文字：" + formula["ocr_text"]
        added |= _add_block(result, text, {**locator, "recognition": "formula_ocr",
                                           "bbox": formula["bbox"], "formula": index,
                                           "label": locator["label"] + f" · 图片公式 {index}"}, heading, "formula")
        result.formula_blocks += 1
    for warning in recognized.warnings:
        if "低置信度图片文字未纳入知识库" in warning or "图片公式输出为空或结构异常" in warning:
            logger.debug("%s: %s", locator["label"], warning)
            continue
        # A failed model download should not produce the same notice on every slide.
        message = warning if "模型" in warning or "环境未安装" in warning else f"{locator['label']}：{warning}"
        if message not in result.warnings:
            result.warnings.append(message)
    return added


def _add_block(
    result: ParseResult,
    text: str | None,
    locator: dict[str, Any],
    heading: str | None = None,
    block_type: str = "text",
) -> bool:
    cleaned = _clean_text(text)
    if not cleaned:
        return False
    result.blocks.append(
        ParsedBlock(cleaned, locator, heading, block_type, len(result.blocks))
    )
    return True


def _pdf_has_images(page: Any) -> bool:
    """Inspect image XObjects, including images nested in a form XObject."""
    visited: set[int] = set()

    def visit(resources: Any) -> bool:
        if resources is None:
            return False
        resources = resources.get_object()
        objects = resources.get("/XObject")
        if objects is None:
            return False
        for reference in objects.get_object().values():
            obj = reference.get_object()
            marker = id(obj)
            if marker in visited:
                continue
            visited.add(marker)
            if obj.get("/Subtype") == "/Image":
                return True
            if obj.get("/Subtype") == "/Form" and visit(obj.get("/Resources")):
                return True
        return False

    return visit(page.get("/Resources"))


def _parse_pdf(path: Path, filename: str, visual=None, image_dir=None) -> ParseResult:
    from pypdf import PdfReader

    result = ParseResult(image_dir=image_dir)
    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as exc:
            raise ValueError("PDF 已加密，无法提取文字。请先解密后重新上传。") from exc
        if not unlocked:
            raise ValueError("PDF 需要密码。请先解密后重新上传。")
        result.warnings.append("PDF 带有加密标记，已使用空密码打开。")
    result.units_total = len(reader.pages)
    result.warnings.append("PDF 保留文字层并补充页面 OCR。")
    rendered = None
    if image_dir is not None or (visual is not None and visual.settings.ocr_enabled):
        try:
            import pypdfium2 as pdfium
            rendered = pdfium.PdfDocument(str(path), password="")
        except Exception:
            result.warnings.append("PDF 页面渲染失败，未进行 OCR，当前仅包含可提取的文字层。")
    try:
        for page_number, page in enumerate(reader.pages, start=1):
            locator = {"filename": filename, "page": page_number, "label": f"第 {page_number} 页"}
            text = ""
            try:
                text = page.extract_text() or ""
            except Exception:
                result.warnings.append(f"第 {page_number} 页文字层提取失败，尝试页面 OCR。")
            with_text = _add_block(result, text, locator)
            if rendered is not None:
                render_page = bitmap = None
                try:
                    render_page = rendered[page_number - 1]
                    width, height = render_page.get_size()
                    max_pixels = visual.settings.ocr_max_image_pixels if visual is not None else 4000000
                    scale = min(2.5, (max_pixels / max(1, width * height)) ** 0.5)
                    bitmap = render_page.render(scale=scale)
                    image = bitmap.to_pil().copy()
                    try:
                        # A page preview also preserves PDF vector drawings. It is
                        # marked separately so it cannot become a new question.
                        with_text |= _add_visual(result, visual, image,
                                                 {**locator, "page_preview": True}, existing_text=text)
                    finally:
                        image.close()
                except Exception:
                    result.warnings.append(f"第 {page_number} 页 OCR 渲染失败，图片内容未纳入。")
                finally:
                    if bitmap is not None:
                        bitmap.close()
                    if render_page is not None:
                        render_page.close()
            elif _pdf_has_images(page):
                result.warnings.append(f"第 {page_number} 页含图片，当前未进行 OCR。")
            if with_text:
                result.units_with_text += 1
            else:
                result.warnings.append(f"第 {page_number} 页未获得可用文字，已跳过。")
    finally:
        if rendered is not None:
            rendered.close()
    return result


def _iter_shapes(shapes: Any, prefix: tuple[int, ...] = ()) -> Iterator[tuple[Any, tuple[int, ...]]]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for index, shape in enumerate(shapes, start=1):
        shape_path = (*prefix, index)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes, shape_path)
        else:
            yield shape, shape_path


def _pptx_table_text(table: Any) -> str:
    # A merged cell appears multiple times in the grid; keep the merge origin once.
    rows = []
    for row in table.rows:
        values = ["" if cell.is_spanned else _clean_text(_pptx_frame_text(cell.text_frame)) for cell in row.cells]
        if any(values):
            rows.append("\t".join(values))
    return "\n".join(rows)


def _pptx_margin_image_key(shape, slide_width, slide_height):
    """Identify small pictures at a slide margin using source and placement."""
    if slide_width <= 0 or slide_height <= 0 or shape.width <= 0 or shape.height <= 0:
        return None
    if shape.width * shape.height > slide_width * slide_height * 0.025:
        return None
    at_margin = (
        shape.top + shape.height <= slide_height * 0.18
        or shape.top >= slide_height * 0.82
        or shape.left + shape.width <= slide_width * 0.12
        or shape.left >= slide_width * 0.88
    )
    if not at_margin:
        return None
    return (
        shape.image.sha1,
        tuple(round(value, 4) for value in (shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom)),
        round(shape.left / slide_width, 3), round(shape.top / slide_height, 3),
        round(shape.width / slide_width, 3), round(shape.height / slide_height, 3),
    )


def _pptx_decoration_keys(presentation):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    occurrences = {}
    for slide_number, slide in enumerate(presentation.slides):
        # Group children use group coordinates, so exclude them from position
        # heuristics rather than mistaking a diagram component for a margin logo.
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            try:
                key = _pptx_margin_image_key(shape, presentation.slide_width, presentation.slide_height)
            except Exception:
                # Leave unreadable pictures to the normal parsing/error path.
                continue
            if key is not None:
                occurrences.setdefault(key, set()).add(slide_number)
    return {key for key, slides in occurrences.items() if len(slides) >= 3}


def _parse_pptx(path: Path, filename: str, visual=None, image_dir=None) -> ParseResult:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(str(path))
    decoration_keys = _pptx_decoration_keys(presentation)
    result = ParseResult(units_total=len(presentation.slides), image_dir=image_dir)
    result.warnings.append("PPTX 读取正文、表格、图片 OCR 及 Office 公式；图表数据、嵌入对象和演讲者备注未纳入。")
    for slide_number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        heading = _clean_text(_pptx_frame_text(title_shape.text_frame)) if title_shape is not None else None
        heading = heading or None
        with_text = False
        unsupported_count = 0
        for shape, shape_path in _iter_shapes(slide.shapes):
            position = ".".join(str(index) for index in shape_path)
            locator = {
                "filename": filename,
                "slide": slide_number,
                "shape": shape_path[-1],
                "shape_path": position,
                "label": f"第 {slide_number} 张幻灯片 · 形状 {position}",
            }
            if shape.has_table:
                with_text |= _add_block(result, _pptx_table_text(shape.table), locator, heading, "table")
            elif shape.has_text_frame:
                block_type = "heading" if title_shape is not None and shape.shape_id == title_shape.shape_id else "text"
                with_text |= _add_block(result, _pptx_frame_text(shape.text_frame), locator, heading, block_type)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    if len(shape_path) == 1 and _pptx_margin_image_key(
                        shape, presentation.slide_width, presentation.slide_height
                    ) in decoration_keys:
                        continue
                    from io import BytesIO
                    from PIL import Image
                    with Image.open(BytesIO(shape.image.blob)) as image:
                        width, height = image.size
                        bounds = (int(width * max(0, shape.crop_left)), int(height * max(0, shape.crop_top)),
                                  int(width * (1 - max(0, shape.crop_right))), int(height * (1 - max(0, shape.crop_bottom))))
                        if bounds[2] <= bounds[0] or bounds[3] <= bounds[1]:
                            raise ValueError("图片裁切区域为空")
                        with image.crop(bounds) as cropped:
                            with_text |= _add_visual(result, visual, cropped, locator, heading)
                except Exception:
                    result.warnings.append(f"{locator['label']}图片格式或裁切不支持，未进行 OCR。")
            elif getattr(shape, "has_chart", False) or shape.shape_type in {
                MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
                MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
            }:
                unsupported_count += 1
            with_text |= _add_office_math(result, shape.element, locator, heading)
        if with_text:
            result.units_with_text += 1
        else:
            result.warnings.append(f"第 {slide_number} 张幻灯片未提取到正文文字，可能为空白或图片资料。")
        if unsupported_count:
            result.warnings.append(f"第 {slide_number} 张幻灯片含 {unsupported_count} 个图表或嵌入对象，未提取其内容。")
    return result


def _docx_table_text(table: Any) -> str:
    # Word returns the same cell object for each position in some merged cells.
    rows = []
    for row in table.rows:
        values = []
        seen_cells: set[int] = set()
        for cell in row.cells:
            marker = id(cell._tc)
            values.append("" if marker in seen_cells else _docx_cell_text(cell))
            seen_cells.add(marker)
        if any(values):
            rows.append("\t".join(values))
    return "\n".join(rows)


def _docx_cell_text(cell: Any) -> str:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    values = []
    for element in cell._tc.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = _clean_text(_docx_paragraph_text(Paragraph(element, cell)))
        elif tag == "tbl":
            text = _docx_table_text(Table(element, cell))
        else:
            continue
        if text:
            values.append(text)
    return "\n".join(values)


def _parse_docx(path: Path, filename: str, visual=None, image_dir=None) -> ParseResult:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    result = ParseResult(image_dir=image_dir)
    result.warnings.append("DOCX 读取正文、表格、正文图片 OCR 及 Office 公式；用段落或表格定位，不提供页码；页眉页脚、文本框及批注未纳入。")
    heading: str | None = None
    paragraph_number = 0
    table_number = 0
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag not in {"p", "tbl"}:
            continue
        result.units_total += 1
        if tag == "p":
            paragraph_number += 1
            paragraph = Paragraph(element, document)
            style = paragraph.style.name if paragraph.style is not None else ""
            is_heading = bool(re.match(r"^(heading|标题)\s*\d*", style, flags=re.IGNORECASE))
            text = _docx_paragraph_text(paragraph)
            if is_heading and _clean_text(text):
                heading = _clean_text(text)
            locator = {"filename": filename, "paragraph": paragraph_number, "label": f"第 {paragraph_number} 段"}
            if heading:
                locator["heading"] = heading
            with_text = _add_block(result, text, locator, heading, "heading" if is_heading else "text")
        else:
            table_number += 1
            locator = {"filename": filename, "table": table_number, "label": f"第 {table_number} 个表格"}
            if heading:
                locator["heading"] = heading
            with_text = _add_block(result, _docx_table_text(Table(element, document)), locator, heading, "table")
        with_text |= _add_office_math(result, element, locator, heading)
        seen = set()
        for item in element.iter():
            if not isinstance(item.tag, str) or item.tag.rsplit("}", 1)[-1] != "blip":
                continue
            relationship = item.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            if relationship in seen:
                continue
            seen.add(relationship)
            image_locator = {**locator, "image": len(seen), "label": locator["label"] + f" · 图片 {len(seen)}"}
            try:
                if not relationship:
                    raise ValueError("外部链接图片不下载")
                with_text |= _add_visual(result, visual, document.part.related_parts[relationship].blob, image_locator, heading)
            except Exception:
                result.warnings.append(f"{image_locator['label']}不是可读取的内嵌图片，未识别。")
        if with_text:
            result.units_with_text += 1
    if result.units_total > result.units_with_text:
        result.warnings.append(f"正文共有 {result.units_total - result.units_with_text} 个空文本段落或表格，未生成检索片段。")
    return result


class DocumentInputError(ValueError):
    """A readable input/configuration problem rather than a corrupt file."""


def _decode_text(path: Path) -> str:
    data = path.read_bytes()
    # Check UTF-32 before UTF-16: their little-endian BOMs share a prefix.
    if data.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        encodings = ("utf-32",)
    elif data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings = ("utf-16",)
    elif data.startswith(b"\xef\xbb\xbf"):
        encodings = ("utf-8-sig",)
    else:
        encodings = ("utf-8", "gb18030")
    for encoding in encodings:
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if "\x00" in text:
            continue
        return text.replace("\r\n", "\n").replace("\r", "\n")
    raise DocumentInputError("无法读取文本编码。请保存为 UTF-8、带 BOM 的 UTF-16 或 GB18030 文本后上传。")


def _parse_text(path: Path, filename: str, visual=None, image_dir=None) -> ParseResult:
    lines = _decode_text(path).splitlines()
    result = ParseResult(units_total=len(lines))
    markdown = Path(filename).suffix.lower() in {".md", ".markdown"}
    heading = None
    pending: list[str] = []
    start_line = 1
    end_line = 1
    pending_size = 0
    fence = None

    def flush():
        nonlocal pending_size
        if pending:
            label = f"第 {start_line} 行" if start_line == end_line else f"第 {start_line}–{end_line} 行"
            locator = {"filename": filename, "line": start_line, "line_end": end_line, "label": label}
            if heading:
                locator["heading"] = heading
            _add_block(result, "\n".join(pending), locator, heading)
            pending.clear()
            pending_size = 0

    for line_number, line in enumerate(lines, 1):
        if line.strip():
            result.units_with_text += 1
        title = re.match(r"^ {0,3}#{1,6}\s+(.+?)\s*#*\s*$", line) if markdown and fence is None else None
        if title:
            flush()
            heading = title.group(1)
            _add_block(result, line, {"filename": filename, "line": line_number,
                                     "line_end": line_number, "heading": heading,
                                     "label": f"第 {line_number} 行"}, heading, "heading")
            continue
        marker = re.match(r"^ {0,3}(`{3,}|~{3,})", line) if markdown else None
        if marker:
            token = marker.group(1)
            if fence is None:
                fence = token
            elif token[0] == fence[0] and len(token) >= len(fence) and not line[marker.end():].strip():
                fence = None
        if not line.strip() and fence is None:
            flush()
            continue
        if len(pending) >= 80 or pending_size + len(line) > 6000:
            flush()
        if not pending:
            start_line = line_number
        end_line = line_number
        pending.append(line)
        pending_size += len(line) + 1
    flush()
    return result


def _parse_image(path: Path, filename: str, visual=None, image_dir=None) -> ParseResult:
    from PIL import Image, ImageOps, ImageSequence

    if image_dir is None and (visual is None or not visual.settings.ocr_enabled):
        raise DocumentInputError("图片资料需要开启 OCR。请在 .env 设置 OCR_ENABLED=true 后重启后端。")
    result = ParseResult(image_dir=image_dir)
    with Image.open(path) as source:
        for page_number, frame in enumerate(ImageSequence.Iterator(source), 1):
            result.units_total += 1
            locator = {"filename": filename, "page": page_number,
                       "label": f"第 {page_number} 页图片"}
            with ImageOps.exif_transpose(frame) as image:
                if _add_visual(result, visual, image, locator):
                    result.units_with_text += 1
    return result


def _xhtml_tag(element) -> str:
    return element.tag.rsplit("}", 1)[-1] if isinstance(element.tag, str) else ""


def _legacy_script_key(text: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", _clean_text(text)))


def _legacy_scripts(manifest: Path) -> dict[str, str]:
    if not manifest.is_file():
        return {}
    if manifest.stat().st_size > 32 * 1024 * 1024:
        raise TikaError("DOC 字符格式结果过大，请拆分文件后重试。")
    try:
        paragraphs = json.loads(manifest.read_text(encoding="utf-8"))
        if not isinstance(paragraphs, list):
            raise ValueError("Expected paragraph list")
        candidates: dict[str, set[str]] = {}
        for paragraph in paragraphs:
            plain, formatted = paragraph["plain"], paragraph["formatted"]
            if not isinstance(plain, str) or not isinstance(formatted, str):
                raise ValueError("Invalid paragraph")
            candidates.setdefault(_legacy_script_key(plain), set()).add(_clean_text(formatted))
    except (ValueError, KeyError, TypeError) as exc:
        raise TikaError("DOC 字符格式结果不完整，请重新导出后上传。") from exc
    # The same plain text can occur in multiple source paragraphs with different
    # scripts; refuse ambiguous replacements rather than copying the wrong one.
    return {key: next(iter(values)) for key, values in candidates.items()
            if key and len(values) == 1 and re.search(r"<(?:sup|sub)>", next(iter(values)))}


def _xhtml_text(element, scripts=None, image_renderer=None) -> str:
    tag = _xhtml_tag(element)
    has_images = any(_xhtml_tag(item) == "img" for item in element.iter())
    if scripts and not (image_renderer and has_images) and tag in {"p", "li", "h1", "h2", "h3", "h4", "h5", "h6"}:
        restored = scripts.get(_legacy_script_key(_xhtml_text(element)))
        if restored is not None:
            return restored + "\n"
    if tag == "img":
        return image_renderer(element) if image_renderer else ""
    if tag == "br":
        return "\n"
    parts = [element.text or ""]
    embedded, hide_alternative = False, False
    for child in element:
        child_tag = _xhtml_tag(child)
        if child_tag == "div" and "embedded" in child.get("class", "").split():
            embedded, hide_alternative = True, False
            value = ""
        elif child_tag == "img":
            value = _xhtml_text(child, scripts, image_renderer)
            hide_alternative = embedded and bool(value)
        elif hide_alternative and child_tag in {"p", "div"}:
            # Tika writes the flattened equation/Visio labels after its preview.
            # The preview is the source of their spatial meaning, not a second stem.
            value = ""
        else:
            value = _xhtml_text(child, scripts, image_renderer)
        tail = child.tail or ""
        parts.extend((value, tail))
        if tail.strip():
            embedded, hide_alternative = False, False
    # Preserve scripts when the source parser provides them. Legacy DOC needs
    # the separate original-run map above because Tika omits these properties.
    if tag in {"sup", "sub"}:
        return _script_text("".join(parts), tag)
    if tag == "span":
        alignment = re.search(r"(?:^|;)\s*vertical-align\s*:\s*(super|sub)\s*(?:;|$)", element.get("style", ""), re.I)
        if alignment:
            return _script_text("".join(parts), "sup" if alignment[1].lower() == "super" else "sub")
    if tag in {"p", "li", "pre", "div", "tr"}:
        parts.append("\n")
    elif tag in {"td", "th"}:
        parts.append("\t")
    content = "".join(parts)
    if image_renderer and has_images and tag == "p" and any(
        _xhtml_tag(child) == "div" and "embedded" in child.get("class", "").split() for child in element
    ):
        # Newlines inserted around embedded previews are XML extraction layout,
        # not separate author paragraphs. Keep formulas in the surrounding sentence.
        content = re.sub(r"\s*\n\s*", " ", content)
    return content


def _xhtml_table_text(element, scripts=None, image_renderer=None) -> str:
    def rows(container):
        for child in container:
            tag = _xhtml_tag(child)
            if tag == "tr":
                yield child
            elif tag != "table":
                yield from rows(child)

    values = []
    for row in rows(element):
        cells = [_clean_text(_xhtml_text(cell, scripts, image_renderer)) for cell in row if _xhtml_tag(cell) in {"td", "th"}]
        if any(cells):
            values.append("\t".join(cells))
    return "\n".join(values)


def _xhtml_units(element):
    tag = _xhtml_tag(element)
    classes = set(element.get("class", "").split())
    # Match the modern Office parsers' body-only scope; do not ingest slide notes or comments.
    if classes & {"notes", "slide-comments", "slide-master-content", "slide-header", "slide-footer", "header", "footer"}:
        return
    if tag == "table":
        yield element, "table"
    elif tag in {"p", "li", "pre", "img", "h1", "h2", "h3", "h4", "h5", "h6"}:
        yield element, "heading" if re.fullmatch(r"h[1-6]", tag) else "image" if tag == "img" else "text"
    else:
        for child in element:
            yield from _xhtml_units(child)


def _parse_tika(path: Path, filename: str, visual=None, image_dir=None) -> ParseResult:
    result = ParseResult(image_dir=image_dir)
    include_images = image_dir is not None or (visual is not None and visual.settings.ocr_enabled)
    with extracted_document(path, include_images=include_images) as (xml, assets):
        scripts = _legacy_scripts(xml.parent / "word-scripts.json")
        try:
            root = ElementTree.parse(xml).getroot()
        except ElementTree.ParseError as exc:
            raise TikaError("旧格式资料未能提取完整内容，请重新导出后上传。") from exc
        body = next((item for item in root.iter() if _xhtml_tag(item) == "body"), None)
        if body is None:
            raise TikaError("旧格式资料没有可读取的正文。")
        files = [item for item in assets.rglob("*") if item.is_file() and item.resolve().is_relative_to(assets.resolve())]
        used_images: set[Path] = set()
        image_markdown = {}
        image_number = 0

        def add_image(image_path, locator, heading, *, inline=False):
            nonlocal image_number
            if not include_images or image_path in used_images:
                return image_markdown.get(image_path, "") if inline else False
            used_images.add(image_path)
            image_number += 1
            try:
                before = len(result.blocks)
                data = image_path.read_bytes()
                if inline:
                    latex = wmf_formula(data)
                    if latex:
                        markdown = "$" + latex + "$"
                        image_markdown[image_path] = markdown
                        result.formula_blocks += 1
                        return markdown
                label = locator["label"] + f" · 图片 {image_number}"
                if inline and data.startswith(b"\x01\x00\x09\x00"):
                    label += " · 行内公式"
                added = _add_visual(result, visual, data, {
                    **locator, "image": image_number,
                    "label": label,
                }, heading)
                if inline:
                    pictures = [block for block in result.blocks[before:] if block.block_type == "image"]
                    if pictures:
                        markdown = pictures[0].text
                        image_markdown[image_path] = markdown
                        result.blocks[before:] = [block for block in result.blocks[before:] if block.block_type != "image"]
                        return markdown
                    return ""
                return added
            except Exception:
                logger.debug("Skipping unreadable Tika image: %s", image_path.name, exc_info=True)
                return "" if inline else False

        def image_path_for(item):
            reference = item.get("src", "")
            if not reference.startswith("embedded:"):
                return None
            name = unquote(reference[len("embedded:"):]).replace("\\", "/")
            basename = name.rsplit("/", 1)[-1]
            matches = [file for file in files if file.relative_to(assets).as_posix() == name or file.name == basename]
            if not matches:
                matches = [file for file in files if file.name.endswith("_" + basename)]
            return matches[0] if len(matches) == 1 else None

        def unit_images(element, locator, heading):
            added = False
            for item in element.iter():
                if _xhtml_tag(item) != "img":
                    continue
                image_path = image_path_for(item)
                if image_path is not None:
                    added |= add_image(image_path, locator, heading)
            return added

        is_ppt = Path(filename).suffix.lower() == ".ppt"
        slides = [item for item in body.iter() if _xhtml_tag(item) == "div" and "slide" in item.get("class", "").split()] if is_ppt else []
        if is_ppt and not slides:
            raise TikaError("未读取到旧版 PPT 的幻灯片结构，请重新导出后上传。")
        containers = slides if is_ppt else [body]
        heading = None
        paragraph_number = table_number = 0
        for slide_number, container in enumerate(containers, 1):
            before_slide = len(result.blocks)
            if is_ppt:
                heading = None
                result.units_total += 1
            for unit_number, (element, kind) in enumerate(_xhtml_units(container), 1):
                text = _xhtml_table_text(element, scripts) if kind == "table" else _xhtml_text(element, scripts)
                if kind == "heading" and _clean_text(text):
                    heading = _clean_text(text)
                if is_ppt:
                    locator = {"filename": filename, "slide": slide_number, "tika_unit": unit_number,
                               "label": f"第 {slide_number} 张幻灯片 · 文本块 {unit_number}"}
                else:
                    result.units_total += 1
                    if kind == "table":
                        table_number += 1
                        locator = {"filename": filename, "table": table_number,
                                   "label": f"第 {table_number} 个表格（提取顺序）"}
                    else:
                        paragraph_number += 1
                        locator = {"filename": filename, "paragraph": paragraph_number,
                                   "label": f"第 {paragraph_number} 段（提取顺序）"}
                if heading:
                    locator["heading"] = heading
                if image_dir is not None:
                    before_visual = len(result.blocks)
                    def render_image(item):
                        image_path = image_path_for(item)
                        return add_image(image_path, locator, heading, inline=True) if image_path is not None else ""
                    text = (_xhtml_table_text(element, scripts, render_image) if kind == "table"
                            else _xhtml_text(element, scripts, render_image))
                    visuals = result.blocks[before_visual:]
                    del result.blocks[before_visual:]
                    added = _add_block(result, text, locator, heading, kind)
                    result.blocks.extend(visuals)
                    for order, block in enumerate(result.blocks):
                        block.order = order
                    added |= bool(visuals)
                else:
                    added = _add_block(result, text, locator, heading, kind) if kind != "image" else False
                    added |= unit_images(element, locator, heading)
                if added and not is_ppt:
                    result.units_with_text += 1
            if is_ppt and len(result.blocks) > before_slide:
                result.units_with_text += 1
        # Older PPTs sometimes expose pictures at presentation level without slide anchors.
        # Keep useful OCR, but do not invent a slide/page number for those pictures.
        for image_path in files:
            if image_path in used_images or image_path.suffix.lower() not in IMAGE_SUFFIXES | {".gif", ".wmf", ".emf"}:
                continue
            locator = {"filename": filename, "label": "嵌入图片（未提供原文位置）"}
            result.units_total += 1
            if add_image(image_path, locator, None):
                result.units_with_text += 1
    return result


def parse_document(path: Path, filename: str, visual=None, *, image_dir=None) -> ParseResult:
    """Read supported documents, text and images or raise a readable error.

    ``units`` means pages/frames, slides, body paragraphs/tables or text lines.
    ``order`` is zero-based and contiguous among non-empty extracted blocks.
    All source positions (page, slide, paragraph, table) are one-based.
    """
    path = Path(path)
    suffix = Path(filename).suffix.lower()
    parsers = {".pdf": _parse_pdf, ".pptx": _parse_pptx, ".docx": _parse_docx}
    parsers.update({extension: _parse_text for extension in TEXT_SUFFIXES})
    parsers.update({extension: _parse_image for extension in IMAGE_SUFFIXES})
    parsers.update({extension: _parse_tika for extension in TIKA_SUFFIXES})
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持该文件格式。请上传 {SUPPORTED_FORMATS_LABEL}。")
    if not path.is_file():
        raise ValueError("资料文件不存在，无法解析。")
    if path.stat().st_size == 0:
        raise ValueError("资料文件为空，无法解析。")
    try:
        if visual is not None:
            visual.begin_document()
        result = parsers[suffix](path, filename, visual, image_dir=image_dir)
    except ImportError as exc:
        raise ValueError(f"解析依赖未安装：{exc.name or '未知依赖'}。请安装后重试。") from exc
    except (TikaError, DocumentInputError):
        raise
    except ValueError as exc:
        if suffix == ".pdf" and ("PDF 已加密" in str(exc) or "PDF 需要密码" in str(exc)):
            raise
        raise ValueError("文件解析失败，可能损坏、格式与扩展名不符或包含不支持的结构。请重新导出后上传。") from exc
    except Exception as exc:
        raise ValueError("文件解析失败，可能损坏、格式与扩展名不符或包含不支持的结构。请重新导出后上传。") from exc
    if not result.blocks:
        for warning in result.warnings:
            logger.debug("%s: %s", filename, warning)
        raise ValueError("未提取到可用内容。")
    # Partial recognition failures are diagnostics, not a user review workflow.
    for warning in result.warnings:
        logger.debug("%s: %s", filename, warning)
    result.warnings.clear()
    return result
